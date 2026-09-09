from __future__ import annotations

import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


RUN_DIR = Path(__file__).resolve().parent
SOURCE_WORKBOOK = RUN_DIR.parent / "execution" / "luna-existing-marker-correctness" / "staged" / "our_markers.xlsx"
OUTPUT_PATH = RUN_DIR / "final-change-set.json"
INVALID_IDS = {"M00157", "M00188", "M00189", "M00209", "M00210", "M00222", "M00223", "M00829"}
UNRESOLVED_IDS = {"M01806", "M01807"}
CORRECTED_IDS = {"M00285", "M01645", "M01808", "M02348", "M02474"}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load_source_rows() -> dict[str, dict[str, object]]:
    workbook = load_workbook(SOURCE_WORKBOOK, read_only=True, data_only=False)
    sheet = workbook["markers"]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value) for value in next(rows)]
    result = {
        str(row[headers.index("marker_id")]): dict(zip(headers, row, strict=True))
        for row in rows
        if row[headers.index("marker_id")]
    }
    workbook.close()
    return result


def load_decisions() -> dict[str, dict[str, object]]:
    decisions: dict[str, dict[str, object]] = {}
    paths = [RUN_DIR / "decisions-context-2922.json", *sorted(RUN_DIR.glob("decisions-review-batch-*.json"))]
    for path in paths:
        document = json.loads(path.read_text(encoding="utf-8"))
        for decision in document["decisions"]:
            marker_id = decision["marker_id"]
            if marker_id in INVALID_IDS | UNRESOLVED_IDS | CORRECTED_IDS:
                decisions[marker_id] = decision
    return decisions


def main() -> None:
    source_rows = load_source_rows()
    decisions = load_decisions()
    target_ids = INVALID_IDS | UNRESOLVED_IDS | CORRECTED_IDS
    missing = sorted(target_ids - decisions.keys())
    if missing:
        raise ValueError(f"Missing final decisions: {missing}")

    baseline_hash = sha256(SOURCE_WORKBOOK)
    changes = []
    for marker_id in sorted(target_ids):
        decision = decisions[marker_id]
        if marker_id in CORRECTED_IDS:
            operation = "update"
            after_values = decision["after_values"]
        elif marker_id in UNRESOLVED_IDS:
            operation = "move_to_unresolved"
            after_values = None
        else:
            operation = "remove"
            after_values = None
        changes.append(
            {
                "change_id": f"ASTRA-FINAL-{operation.upper()}-{marker_id}",
                "decision_id": decision["decision_id"],
                "operation": operation,
                "marker_id": marker_id,
                "before_values": source_rows[marker_id],
                "after_values": after_values,
                "evidence_ids": [],
                "reason": decision["astra_reason"],
                "expected_baseline_sha256": baseline_hash,
                "executor": "GPT-6 Astra",
                "decided_at": "2026-09-09",
            }
        )

    output = {
        "run_id": "20260909-astra-final-marker-publication",
        "source_workbook": str(SOURCE_WORKBOOK),
        "source_sha256": baseline_hash,
        "change_count": len(changes),
        "operation_counts": {"update": 5, "remove": 8, "move_to_unresolved": 2},
        "expected_published_marker_count": 2480,
        "changes": changes,
    }
    OUTPUT_PATH.write_text(json.dumps(output, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(OUTPUT_PATH), "changes": len(changes), "expected_published": 2480}, ensure_ascii=False))


if __name__ == "__main__":
    main()
