from __future__ import annotations

import hashlib
import json
from datetime import datetime
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


RUN_DIR = Path(__file__).resolve().parent
MAIN_PATH = RUN_DIR / "staged" / "our_markers_final.xlsx"
BY_CELL_PATH = RUN_DIR / "staged" / "our_markers_by_cell_final.xlsx"
REVIEWING_PATH = RUN_DIR / "staged" / "our_markers_reviewing.xlsx"
OUTPUT_PATH = RUN_DIR / "final-acceptance-validation.json"
INVALID_IDS = {"M00157", "M00188", "M00189", "M00209", "M00210", "M00222", "M00223", "M00829"}
UNRESOLVED_IDS = {"M01806", "M01807"}
CORRECTED_IDS = {"M00285", "M01645", "M01808", "M02348", "M02474"}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def records(sheet: Any, id_header: str | None = None) -> tuple[list[str], list[dict[str, Any]]]:
    values = sheet.iter_rows(values_only=True)
    headers = [str(value) if value is not None else "" for value in next(values)]
    output = []
    for row in values:
        item = dict(zip(headers, row, strict=False))
        if id_header is None:
            if any(value not in (None, "") for value in row):
                output.append(item)
        elif item.get(id_header):
            output.append(item)
    return headers, output


def marker_display(record: dict[str, Any]) -> str:
    polarity = str(record.get("marker_polarity") or "").lower()
    if polarity == "negative":
        return f"{record['gene_symbol']}（阴性）"
    if polarity == "low":
        return f"{record['gene_symbol']}（低表达）"
    if polarity in {"unknown", ""}:
        return f"{record['gene_symbol']}（方向未明）"
    return str(record["gene_symbol"])


def normalized_doi(paper_id: str) -> str:
    return paper_id[4:].replace("_", "/") if paper_id.startswith("DOI_") else paper_id


def error_cells(workbook: Any) -> list[str]:
    errors = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "e":
                    errors.append(f"{sheet.title}!{cell.coordinate}:{cell.value}")
    return errors


def main() -> None:
    main_workbook = load_workbook(MAIN_PATH, read_only=False, data_only=False)
    headers, main_records = records(main_workbook["markers"], "marker_id")
    main_by_id = {str(record["marker_id"]): record for record in main_records}
    unresolved_headers, unresolved_records = records(main_workbook["unresolved_hold_20260909"], "marker_id")
    unresolved_by_id = {str(record["marker_id"]): record for record in unresolved_records}
    table_refs = {table.name: table.ref for table in main_workbook["markers"].tables.values()}
    main_errors = error_cells(main_workbook)

    reviewing_workbook = load_workbook(REVIEWING_PATH, read_only=True, data_only=False)
    _, reviewing_records = records(reviewing_workbook["markers"], "marker_id")
    reviewing_by_id = {str(record["marker_id"]): record for record in reviewing_records}
    reviewing_workbook.close()

    deferred_decisions = []
    for decision_path in sorted(RUN_DIR.glob("decisions-deferred-batch-*.json")):
        deferred_decisions.extend(json.loads(decision_path.read_text(encoding="utf-8"))["decisions"])
    deferred_accepted_ids = {decision["marker_id"] for decision in deferred_decisions if decision["publication_eligible"]}

    by_cell_workbook = load_workbook(BY_CELL_PATH, read_only=False, data_only=False)
    _, by_cell_records = records(by_cell_workbook["Marker表"])
    _, panel_records = records(by_cell_workbook["图表索引"])
    info = {str(row[0]): row[1] for row in by_cell_workbook["说明"].iter_rows(values_only=True) if row[0] is not None}
    by_cell_errors = error_cells(by_cell_workbook)

    expected_groups: dict[tuple[str, str, str], set[str]] = defaultdict(set)
    for record in main_records:
        author_label = str(record["cell_type"])
        if record.get("subtype"):
            author_label = f"{author_label}（{record['subtype']}）"
        expected_groups[(normalized_doi(str(record["paper_id"])), author_label, str(record["species"]))].add(marker_display(record))
    def species_from_basis(value: Any) -> str:
        text = str(value or "")
        prefix = "物种："
        if prefix not in text:
            return ""
        return text.split(prefix, 1)[1].split("；", 1)[0]
    actual_groups = {
        (str(record["DOI"]), str(record["作者细胞名称"]), species_from_basis(record.get("依据"))): set(str(record["Marker"]).split("；"))
        for record in by_cell_records
    }
    group_keys_missing = sorted(set(expected_groups) - set(actual_groups))
    group_keys_extra = sorted(set(actual_groups) - set(expected_groups))
    group_marker_mismatches = {
        " | ".join(key): {
            "expected": sorted(expected_groups[key]),
            "actual": sorted(actual_groups[key]),
        }
        for key in sorted(set(expected_groups) & set(actual_groups))
        if expected_groups[key] != actual_groups[key]
    }

    correction_mismatches = {}
    for marker_id in CORRECTED_IDS:
        decision = None
        for path in sorted(RUN_DIR.glob("decisions-review-batch-*.json")):
            matches = [item for item in json.loads(path.read_text(encoding="utf-8"))["decisions"] if item["marker_id"] == marker_id]
            if matches:
                decision = matches[0]
                break
        expected = decision["after_values"] if decision else None
        if expected is None or any(main_by_id[marker_id].get(header) != expected.get(header) for header in headers):
            correction_mismatches[marker_id] = True

    def equivalent_excel_value(left: Any, right: Any) -> bool:
        if left == right:
            return True
        if isinstance(left, datetime) and isinstance(right, (int, float)):
            return (left - datetime(1899, 12, 30)).days == int(right)
        if isinstance(right, datetime) and isinstance(left, (int, float)):
            return (right - datetime(1899, 12, 30)).days == int(left)
        return False

    unresolved_snapshot_mismatches = []
    for marker_id in UNRESOLVED_IDS:
        if any(not equivalent_excel_value(unresolved_by_id[marker_id].get(header), reviewing_by_id[marker_id].get(header)) for header in headers):
            unresolved_snapshot_mismatches.append(marker_id)

    checks = {
        "published_count_is_2480": len(main_records) == 2480,
        "published_ids_unique": len(main_by_id) == len(main_records),
        "invalid_absent_from_published": not (INVALID_IDS & main_by_id.keys()),
        "unresolved_absent_from_published": not (UNRESOLVED_IDS & main_by_id.keys()),
        "unresolved_sheet_exact": set(unresolved_by_id) == UNRESOLVED_IDS,
        "unresolved_snapshots_preserved": not unresolved_snapshot_mismatches,
        "all_397_deferred_accepts_published": deferred_accepted_ids <= main_by_id.keys() and len(deferred_accepted_ids) == 397,
        "five_corrections_applied": not correction_mismatches,
        "marker_table_ref_exact": table_refs.get("OurMarkersTable") == "A1:AE2481",
        "by_cell_group_keys_exact": not group_keys_missing and not group_keys_extra,
        "by_cell_marker_sets_exact": not group_marker_mismatches,
        "by_cell_reported_total_is_2480": info.get("Marker总数") == 2480,
        "no_formula_or_cell_errors": not main_errors and not by_cell_errors,
    }
    output = {
        "run_id": "20260909-astra-final-acceptance",
        "status": "passed" if all(checks.values()) else "failed",
        "checks": checks,
        "counts": {
            "published_markers": len(main_records),
            "published_unique_ids": len(main_by_id),
            "deferred_accepted_published": len(deferred_accepted_ids & main_by_id.keys()),
            "unresolved_hold": len(unresolved_records),
            "by_cell_groups": len(by_cell_records),
            "panel_index_rows": len(panel_records),
            "papers": len({record["paper_id"] for record in main_records}),
        },
        "diagnostics": {
            "table_refs": table_refs,
            "unresolved_snapshot_mismatches": unresolved_snapshot_mismatches,
            "correction_mismatches": correction_mismatches,
            "group_keys_missing": group_keys_missing,
            "group_keys_extra": group_keys_extra,
            "group_marker_mismatches": group_marker_mismatches,
            "main_errors": main_errors,
            "by_cell_errors": by_cell_errors,
        },
        "hashes": {"our_markers_final.xlsx": sha256(MAIN_PATH), "our_markers_by_cell_final.xlsx": sha256(BY_CELL_PATH)},
    }
    OUTPUT_PATH.write_text(json.dumps(output, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
    print(json.dumps(output, ensure_ascii=False, indent=2, default=str))
    if output["status"] != "passed":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
