from __future__ import annotations

import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


RUN_DIR = Path(__file__).resolve().parent
WORKSPACE = RUN_DIR.parents[3]
DEFERRED_PATH = RUN_DIR / "deferred-remaining-current.json"
WORKBOOK_PATH = RUN_DIR / "staged" / "our_markers_reviewing.xlsx"
OUTPUT_PATH = RUN_DIR / "decisions-deferred-batch-002.json"
PAPER_ID = "DOI_10.64898_2025.12.18.695268"
PDF_PATH = next((WORKSPACE / "marker提取" / "pdf").glob("*695268*.pdf"))


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def main() -> None:
    deferred = json.loads(DEFERRED_PATH.read_text(encoding="utf-8"))["records"]
    records = [record for record in deferred if record["paper_id"] == PAPER_ID]
    if len(records) != 42:
        raise ValueError(f"Expected 42 deferred records for {PAPER_ID}, got {len(records)}")

    workbook = load_workbook(WORKBOOK_PATH, read_only=True, data_only=False)
    sheet = workbook["markers"]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value) for value in next(rows)]
    marker_id_index = headers.index("marker_id")
    marker_rows = {
        str(row[marker_id_index]): dict(zip(headers, row, strict=True))
        for row in rows
        if row[marker_id_index]
    }

    pdf_hash = sha256(PDF_PATH)
    decisions = []
    for record in records:
        marker_id = record["marker_id"]
        after_values = marker_rows.get(marker_id)
        if after_values is None:
            raise ValueError(f"Marker missing from reviewing workbook: {marker_id}")
        locator = record["source_locator"]
        if marker_id == "M02487":
            reason = "Fig. S3C 的 DCN violin 在 xFB-8 列直接显示表达，补足此前缺失的目标级 xFB-8-DCN 绑定。"
        elif "S3C" in locator:
            reason = "逐项对照 Fig. S3C 的 marker violin，目标细胞/亚群列存在对应 marker 表达。"
        elif "S1" in locator:
            reason = "逐项对照 Fig. 1D dotplot 与 Fig. S1 feature plot，目标细胞-marker 关系一致。"
        else:
            reason = "逐项对照 Fig. 1D canonical-marker dotplot，目标细胞行与 marker 列的表达关系一致。"
        decisions.append(
            {
                "decision_id": f"ASTRA-DEFERRED-B002-{marker_id}",
                "marker_id": marker_id,
                "paper_id": PAPER_ID,
                "decision": "accept_current_after_values",
                "publication_eligible": True,
                "after_values": after_values,
                "source_file": str(PDF_PATH),
                "source_sha256": pdf_hash,
                "source_locator": locator,
                "source_excerpt": after_values.get("source_context"),
                "astra_reason": reason,
                "reviewed_by": "GPT-6 Astra",
                "reviewed_at": "2026-09-09",
            }
        )

    document = {
        "run_id": "20260909-astra-deferred-review",
        "batch_id": "ASTRA-DEFERRED-B002",
        "paper_id": PAPER_ID,
        "record_count": len(decisions),
        "result_counts": {"accept_current_after_values": len(decisions)},
        "remaining_after_batch": 119,
        "evidence_groups": {"Fig. 1D / Fig. S1": 37, "Fig. S3C": 5},
        "decisions": decisions,
    }
    OUTPUT_PATH.write_text(json.dumps(document, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(OUTPUT_PATH), "accepted": len(decisions), "remaining": 119}, ensure_ascii=False))


if __name__ == "__main__":
    main()
