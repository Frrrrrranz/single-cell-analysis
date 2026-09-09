from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
RUN = ROOT / "execution" / "luna-existing-marker-correctness"
OUTPUT = Path(__file__).with_name("workbook-data-audit.json")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def sheet_records(path: Path, sheet_name: str, id_column: str) -> tuple[list[str], dict[str, dict[str, Any]]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    sheet = workbook[sheet_name]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value) if value is not None else "" for value in next(rows)]
    records: dict[str, dict[str, Any]] = {}
    for values in rows:
        record = dict(zip(headers, values))
        record_id = record.get(id_column)
        if record_id:
            records[str(record_id)] = record
    workbook.close()
    return headers, records


def changed_fields(before: dict[str, Any], after: dict[str, Any], headers: list[str]) -> dict[str, dict[str, Any]]:
    return {
        header: {"before": before.get(header), "after": after.get(header)}
        for header in headers
        if before.get(header) != after.get(header)
    }


def main() -> None:
    baseline_main = RUN / "baseline" / "our_markers.xlsx"
    staged_main = RUN / "staged" / "our_markers.xlsx"
    staged_by_cell = RUN / "staged" / "our_markers_by_cell.xlsx"
    historical_by_cell = ROOT / "execution" / "luna-single-model" / "baseline" / "20260906-luna-single-model" / "our_markers_by_cell.xlsx"
    headers, before = sheet_records(baseline_main, "markers", "marker_id")
    staged_headers, after = sheet_records(staged_main, "markers", "marker_id")
    assert headers == staged_headers

    added_ids = sorted(set(after) - set(before))
    removed_ids = sorted(set(before) - set(after))
    common_diffs = {
        marker_id: changed_fields(before[marker_id], after[marker_id], headers)
        for marker_id in sorted(set(before) & set(after))
    }
    common_diffs = {marker_id: fields for marker_id, fields in common_diffs.items() if fields}

    reviews = json.loads((RUN / "review-records.json").read_text(encoding="utf-8"))["records"]
    expected_update_ids = sorted(record["marker_id"] for record in reviews if record.get("action") == "update")
    expected_restore_ids = sorted(record["marker_id"] for record in reviews if record.get("action") == "restore")

    workbook = load_workbook(staged_main, read_only=True, data_only=False)
    archive = workbook["audit_archive_20260907"]
    archive_headers = [cell.value for cell in next(archive.iter_rows(min_row=1, max_row=1))]
    archive_row_count = sum(1 for row in archive.iter_rows(min_row=2, values_only=True) if any(value is not None for value in row))
    workbook.close()

    by_cell_workbook = load_workbook(staged_by_cell, read_only=True, data_only=True)
    marker_sheet = by_cell_workbook["Marker表"]
    marker_rows = sum(1 for row in marker_sheet.iter_rows(min_row=2, values_only=True) if any(value is not None for value in row))
    index_sheet = by_cell_workbook["图表索引"]
    index_rows = sum(1 for row in index_sheet.iter_rows(min_row=2, values_only=True) if any(value is not None for value in row))
    info_sheet = by_cell_workbook["说明"]
    info = {str(row[0]): row[1] for row in info_sheet.iter_rows(values_only=True) if row[0] is not None}
    by_cell_workbook.close()

    historical_workbook = load_workbook(historical_by_cell, read_only=True, data_only=True)
    historical_marker = historical_workbook["Marker表"]
    historical_rows = list(historical_marker.iter_rows(min_row=2, values_only=True))
    historical_index = historical_workbook["图表索引"]
    historical_index_rows = list(historical_index.iter_rows(min_row=2, values_only=True))
    historical_workbook.close()

    staged_workbook = load_workbook(staged_by_cell, read_only=True, data_only=True)
    staged_marker_rows = list(staged_workbook["Marker表"].iter_rows(min_row=2, values_only=True))
    staged_index_rows = list(staged_workbook["图表索引"].iter_rows(min_row=2, values_only=True))
    staged_workbook.close()

    def nonempty_by_column(rows: list[tuple[Any, ...]], indexes: list[int]) -> dict[str, int]:
        return {str(index + 1): sum(row[index] not in (None, "") for row in rows) for index in indexes}

    result = {
        "run_id": "20260908-astra-audit-after-luna-existing-marker-correctness",
        "hashes": {
            "baseline_main": sha256(baseline_main),
            "staged_main": sha256(staged_main),
            "staged_by_cell": sha256(staged_by_cell),
        },
        "main_workbook": {
            "baseline_ids": len(before),
            "staged_ids": len(after),
            "added_ids": added_ids,
            "removed_ids": removed_ids,
            "common_ids_with_field_changes": len(common_diffs),
            "changed_ids": common_diffs,
            "expected_update_ids": expected_update_ids,
            "expected_restore_ids": expected_restore_ids,
            "unexpected_changed_ids": sorted(set(common_diffs) - set(expected_update_ids)),
            "expected_updates_missing_from_diff": sorted(set(expected_update_ids) - set(common_diffs)),
        },
        "archive": {
            "headers": archive_headers,
            "header_count": len(archive_headers),
            "data_rows": archive_row_count,
            "contains_all_source_headers": all(header in archive_headers for header in headers),
        },
        "by_cell_workbook": {
            "marker_rows": marker_rows,
            "index_rows": index_rows,
            "reported_data_source": info.get("数据来源"),
            "reported_marker_total": info.get("Marker总数"),
            "reported_panel_author_label_rows": info.get("面板—作者标签记录数"),
            "reported_paper_count": info.get("论文数"),
            "historical_marker_metadata_nonempty_C_to_F": nonempty_by_column(historical_rows, [2, 3, 4, 5]),
            "staged_marker_metadata_nonempty_C_to_F": nonempty_by_column(staged_marker_rows, [2, 3, 4, 5]),
            "historical_index_metadata_nonempty_C_to_F": nonempty_by_column(historical_index_rows, [2, 3, 4, 5]),
            "staged_index_metadata_nonempty_C_to_F": nonempty_by_column(staged_index_rows, [2, 3, 4, 5]),
        },
    }
    OUTPUT.write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2, default=str))


if __name__ == "__main__":
    main()
